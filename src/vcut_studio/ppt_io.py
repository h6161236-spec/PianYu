from __future__ import annotations

import math
import re
import shutil
import subprocess
import tempfile
import textwrap
from dataclasses import dataclass
from pathlib import Path
from uuid import uuid4
from collections.abc import Callable

from .models import SlidePage
from .process_utils import subprocess_windowless_kwargs


@dataclass(frozen=True, slots=True)
class PptImportResult:
    slides: list[SlidePage]
    rendered_slides_dir: Path
    deck_summary: str = ""


def _run_command(command: list[str], error_message: str) -> None:
    completed = subprocess.run(
        command,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
        **subprocess_windowless_kwargs(),
    )
    if completed.returncode == 0:
        return
    stderr_lines = (completed.stderr or completed.stdout or "").strip().splitlines()
    stderr_tail = "\n".join(stderr_lines[-12:])
    raise RuntimeError(stderr_tail or error_message)


def _powershell_utf8_expression(value: str) -> str:
    import base64

    encoded = base64.b64encode(value.encode("utf-8")).decode("ascii")
    return "[System.Text.Encoding]::UTF8.GetString([System.Convert]::FromBase64String('" + encoded + "'))"


def _run_powershell(script: str, error_message: str) -> None:
    import base64

    encoded_script = base64.b64encode(script.encode("utf-16-le")).decode("ascii")
    _run_command(
        ["powershell", "-NoProfile", "-EncodedCommand", encoded_script],
        error_message,
    )


def _normalize_text(value: object) -> str:
    text = str(value or "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [" ".join(line.split()) for line in text.split("\n")]
    return "\n".join(line for line in lines if line).strip()


def _slide_number_from_name(path: Path) -> int:
    match = re.search(r"(\d+)", path.stem)
    if match is None:
        return 0
    return int(match.group(1))


def _find_libreoffice_executable() -> str | None:
    candidates = [
        shutil.which("soffice"),
        shutil.which("soffice.exe"),
        shutil.which("soffice.com"),
        str(Path("C:/Program Files/LibreOffice/program/soffice.exe")),
        str(Path("C:/Program Files (x86)/LibreOffice/program/soffice.exe")),
    ]
    for candidate in candidates:
        if not candidate:
            continue
        if Path(candidate).exists():
            return candidate
    return None


def _convert_legacy_ppt_with_powerpoint(source_path: Path, target_path: Path) -> None:
    target_path.parent.mkdir(parents=True, exist_ok=True)
    script = (
        f"$inputPath = {_powershell_utf8_expression(str(source_path.resolve()))}\n"
        f"$outputPath = {_powershell_utf8_expression(str(target_path.resolve()))}\n"
        "$app = New-Object -ComObject PowerPoint.Application\n"
        "$presentation = $null\n"
        "try {\n"
        "    $presentation = $app.Presentations.Open($inputPath, $false, $false, $false)\n"
        "    $presentation.SaveAs($outputPath, 24)\n"
        "} finally {\n"
        "    if ($presentation -ne $null) { $presentation.Close() }\n"
        "    $app.Quit()\n"
        "}\n"
    )
    _run_powershell(script, "Failed to convert the legacy .ppt file with Microsoft PowerPoint.")


def _convert_legacy_ppt_with_libreoffice(source_path: Path, target_path: Path) -> None:
    executable = _find_libreoffice_executable()
    if executable is None:
        raise RuntimeError("LibreOffice is not installed.")
    target_path.parent.mkdir(parents=True, exist_ok=True)
    _run_command(
        [
            executable,
            "--headless",
            "--convert-to",
            "pptx",
            "--outdir",
            str(target_path.parent),
            str(source_path),
        ],
        "Failed to convert the legacy .ppt file with LibreOffice.",
    )
    converted_path = target_path.parent / f"{source_path.stem}.pptx"
    if converted_path != target_path and converted_path.exists():
        if target_path.exists():
            target_path.unlink()
        converted_path.replace(target_path)


def _prepare_source_for_parsing(source_path: Path, rendered_dir: Path) -> Path:
    if source_path.suffix.lower() != ".ppt":
        return source_path

    converted_dir = rendered_dir / "_converted"
    converted_path = converted_dir / f"{source_path.stem}.pptx"
    conversion_errors: list[str] = []
    for converter in (_convert_legacy_ppt_with_powerpoint, _convert_legacy_ppt_with_libreoffice):
        try:
            converter(source_path, converted_path)
        except Exception as exc:
            conversion_errors.append(str(exc))
            continue
        if converted_path.exists():
            return converted_path
    raise RuntimeError(
        "This .ppt file needs to be converted to .pptx before import, "
        "but both Microsoft PowerPoint and LibreOffice conversion failed.\n"
        + "\n".join(conversion_errors)
    )


def _render_slide_images_with_powerpoint(
    source_path: Path,
    rendered_dir: Path,
    slide_count: int,
    aspect_ratio: float,
) -> dict[int, Path]:
    export_width = 1920
    export_height = max(1080, int(round(export_width / max(0.2, float(aspect_ratio or (16 / 9))))))
    with tempfile.TemporaryDirectory(prefix="vcut_ppt_render_") as temp_dir_name:
        export_dir = Path(temp_dir_name)
        script = (
            f"$inputPath = {_powershell_utf8_expression(str(source_path.resolve()))}\n"
            f"$outputDir = {_powershell_utf8_expression(str(export_dir.resolve()))}\n"
            "$app = New-Object -ComObject PowerPoint.Application\n"
            "$presentation = $null\n"
            "try {\n"
            "    $presentation = $app.Presentations.Open($inputPath, $false, $false, $false)\n"
            f"    $presentation.Export($outputDir, 'PNG', {export_width}, {export_height})\n"
            "} finally {\n"
            "    if ($presentation -ne $null) { $presentation.Close() }\n"
            "    $app.Quit()\n"
            "}\n"
        )
        _run_powershell(script, "Failed to render real slide previews with Microsoft PowerPoint.")

        exported_paths = sorted(
            [path for path in export_dir.iterdir() if path.suffix.lower() == ".png"],
            key=_slide_number_from_name,
        )
        if len(exported_paths) < slide_count:
            raise RuntimeError("PowerPoint exported fewer slide images than expected.")

        preview_paths: dict[int, Path] = {}
        for slide_index, exported_path in enumerate(exported_paths[:slide_count], start=1):
            target_path = rendered_dir / f"slide_{slide_index:03d}.png"
            shutil.copy2(exported_path, target_path)
            preview_paths[slide_index] = target_path
        return preview_paths


def _append_unique_text(buffer: list[str], value: str) -> None:
    normalized = _normalize_text(value)
    if not normalized:
        return
    if normalized in buffer:
        return
    buffer.append(normalized)


def _extract_shape_text(shape: object, buffer: list[str]) -> None:
    if bool(getattr(shape, "has_text_frame", False)):
        _append_unique_text(buffer, getattr(shape, "text", ""))

    if bool(getattr(shape, "has_table", False)):
        table = getattr(shape, "table", None)
        rows = getattr(table, "rows", []) if table is not None else []
        for row in rows:
            for cell in getattr(row, "cells", []):
                _append_unique_text(buffer, getattr(cell, "text", ""))

    nested_shapes = getattr(shape, "shapes", None)
    if nested_shapes is None:
        return
    try:
        for child_shape in nested_shapes:
            _extract_shape_text(child_shape, buffer)
    except TypeError:
        return


def _shape_collection_title(shapes: object) -> str:
    title_shape = getattr(shapes, "title", None)
    if title_shape is None:
        return ""
    return _normalize_text(getattr(title_shape, "text", ""))


def extract_slide_content(slide: object) -> tuple[str, str]:
    title = _shape_collection_title(getattr(slide, "shapes", None))
    collected: list[str] = []
    for shape in getattr(slide, "shapes", []):
        _extract_shape_text(shape, collected)
    if title:
        collected = [item for item in collected if item != title]
    return (title, "\n\n".join(collected).strip())


def extract_slide_notes(slide: object) -> str:
    try:
        notes_slide = slide.notes_slide
    except Exception:
        return ""

    notes_text_frame = getattr(notes_slide, "notes_text_frame", None)
    if notes_text_frame is not None:
        text_value = _normalize_text(getattr(notes_text_frame, "text", ""))
        if text_value and not text_value.isdigit():
            return text_value

    collected: list[str] = []
    for shape in getattr(notes_slide, "shapes", []):
        text_value = _normalize_text(getattr(shape, "text", ""))
        if not text_value or text_value.isdigit():
            continue
        _append_unique_text(collected, text_value)
    return "\n\n".join(collected).strip()


def _best_preview_font(image_font_module: object, font_size: int) -> object:
    font_candidates = [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/msyh.ttf"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in font_candidates:
        if not candidate.exists():
            continue
        try:
            return image_font_module.truetype(str(candidate), font_size)
        except Exception:
            continue
    return image_font_module.load_default()


def _wrap_text(value: str, width: int) -> list[str]:
    normalized = _normalize_text(value)
    if not normalized:
        return []
    lines: list[str] = []
    for paragraph in normalized.split("\n"):
        wrapped = textwrap.wrap(paragraph, width=max(12, width), break_long_words=True)
        if wrapped:
            lines.extend(wrapped)
        else:
            lines.append("")
    return lines


def render_slide_preview_image(
    title: str,
    source_text: str,
    notes_text: str,
    output_path: str | Path,
    *,
    aspect_ratio: float = 16 / 9,
) -> Path:
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        return Path("")

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    width = 1280
    height = max(720, int(round(width / max(0.2, float(aspect_ratio or (16 / 9))))))
    image = Image.new("RGB", (width, height), "#F8FAFC")
    draw = ImageDraw.Draw(image)

    title_font = _best_preview_font(ImageFont, 40)
    body_font = _best_preview_font(ImageFont, 24)
    label_font = _best_preview_font(ImageFont, 20)

    margin_x = 56
    margin_y = 48
    content_width_chars = 64
    cursor_y = margin_y

    draw.rounded_rectangle((28, 28, width - 28, height - 28), radius=24, outline="#CBD5E1", width=2)
    draw.text((margin_x, cursor_y), title or "未命名页面", font=title_font, fill="#0F172A")
    cursor_y += 74

    sections = [
        ("页面文字", source_text),
        ("演讲备注", notes_text),
    ]
    for label, value in sections:
        if not _normalize_text(value):
            continue
        draw.text((margin_x, cursor_y), label, font=label_font, fill="#0F766E")
        cursor_y += 34
        for line in _wrap_text(value, content_width_chars):
            draw.text((margin_x, cursor_y), line, font=body_font, fill="#1E293B")
            cursor_y += 31
            if cursor_y >= height - 70:
                ellipsis = "..."
                draw.text((margin_x, cursor_y), ellipsis, font=body_font, fill="#475569")
                image.save(output_path)
                return output_path
        cursor_y += 18

    if cursor_y <= margin_y + 90:
        draw.text(
            (margin_x, cursor_y),
            "这一页没有可提取的文字内容。",
            font=body_font,
            fill="#475569",
        )

    image.save(output_path)
    return output_path


def import_presentation(
    ppt_path: str | Path,
    rendered_slides_dir: str | Path,
    progress_callback: Callable[[int, str], None] | None = None,
) -> PptImportResult:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError(
            "python-pptx is not installed. Run `pip install python-pptx Pillow` first."
        ) from exc

    source_path = Path(ppt_path)
    if not source_path.exists():
        raise FileNotFoundError(f"PPT file does not exist: {source_path}")

    rendered_dir = Path(rendered_slides_dir)
    rendered_dir.mkdir(parents=True, exist_ok=True)
    if progress_callback is not None:
        progress_callback(5, "正在准备 PPT 文件...")
    parse_source_path = _prepare_source_for_parsing(source_path, rendered_dir)

    if progress_callback is not None:
        progress_callback(12, "正在读取 PPT 结构...")
    presentation = Presentation(str(parse_source_path))
    slide_items = list(presentation.slides)
    slide_width = int(getattr(presentation, "slide_width", 0) or 0)
    slide_height = int(getattr(presentation, "slide_height", 0) or 0)
    aspect_ratio = (slide_width / slide_height) if slide_width > 0 and slide_height > 0 else (16 / 9)
    powerpoint_preview_paths: dict[int, Path] = {}
    try:
        if progress_callback is not None:
            progress_callback(24, "正在渲染页面预览，这一步遇到大文件会稍等一会儿...")
        powerpoint_preview_paths = _render_slide_images_with_powerpoint(
            source_path,
            rendered_dir,
            len(slide_items),
            aspect_ratio,
        )
    except Exception:
        powerpoint_preview_paths = {}

    slides: list[SlidePage] = []
    total_slides = max(1, len(slide_items))
    for slide_index, slide in enumerate(slide_items, start=1):
        if progress_callback is not None:
            progress = 35 + int((slide_index - 1) * 60 / total_slides)
            progress_callback(progress, f"正在提取第 {slide_index}/{total_slides} 页内容...")
        title, source_text = extract_slide_content(slide)
        notes_text = extract_slide_notes(slide)
        preview_path = powerpoint_preview_paths.get(slide_index)
        if preview_path is None or not preview_path.exists():
            preview_path = render_slide_preview_image(
                title=title,
                source_text=source_text,
                notes_text=notes_text,
                output_path=rendered_dir / f"slide_{slide_index:03d}.png",
                aspect_ratio=aspect_ratio,
            )
        estimated_duration_ms = max(
            8_000,
            min(
                90_000,
                int(
                    math.ceil(
                        len((notes_text or source_text or title).split()) / 2.5
                    )
                    * 1000
                ),
            ),
        )
        slides.append(
            SlidePage(
                slide_id=uuid4().hex,
                slide_index=slide_index,
                title=title,
                source_text=source_text,
                notes_text=notes_text,
                preview_image_path=str(preview_path),
                estimated_duration_ms=estimated_duration_ms,
            )
        )

    deck_summary = " / ".join(slide.title for slide in slides[:3] if slide.title).strip()
    if progress_callback is not None:
        progress_callback(100, f"PPT 导入完成，共 {len(slides)} 页。")
    return PptImportResult(
        slides=slides,
        rendered_slides_dir=rendered_dir,
        deck_summary=deck_summary,
    )
